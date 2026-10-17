import json

from docx import Document

from app.db import SessionLocal, init_db
from app.models import ProfileItemStatus
from app.profile_builder.coverage import ledger_summary, record_item
from app.profile_builder.cv_parser import extract_text, ingest_cv
from tests.fakes import FakeLLMProvider


def setup_module():
    init_db()


def test_extract_text_from_docx(tmp_path):
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Hasan Kalin")
    doc.add_paragraph("Backend engineer with Python and SQL experience.")
    doc.save(path)

    text = extract_text(str(path))
    assert "Hasan Kalin" in text
    assert "Python" in text


def test_ingest_cv_records_analyzed_item(tmp_path):
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Backend engineer with Python experience.")
    doc.save(path)

    llm = FakeLLMProvider(
        json.dumps(
            {
                "purpose": "Professional CV",
                "stack": ["Python", "SQL"],
                "role": "Backend Engineer",
                "highlights": ["5 years experience"],
            }
        )
    )
    db = SessionLocal()
    try:
        item = ingest_cv(db, llm, str(path))
        assert item.status == ProfileItemStatus.ANALYZED
        assert item.extracted["stack"] == ["Python", "SQL"]
        assert len(llm.calls) == 1
    finally:
        db.close()


def test_ingest_cv_flags_needs_review_on_malformed_llm_output(tmp_path):
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Some text")
    doc.save(path)

    llm = FakeLLMProvider("not valid json")
    db = SessionLocal()
    try:
        item = ingest_cv(db, llm, str(path))
        assert item.status == ProfileItemStatus.NEEDS_REVIEW
    finally:
        db.close()


def test_ingest_cv_survives_corrupt_pdf_instead_of_crashing(tmp_path):
    """Regression: a corrupt/unreadable file must become a needs_review item,
    not an unhandled exception that 500s the request."""
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"not a real pdf at all")

    llm = FakeLLMProvider("should never be called")
    db = SessionLocal()
    try:
        item = ingest_cv(db, llm, str(path))
        assert item.status == ProfileItemStatus.NEEDS_REVIEW
        assert "parse_error" in item.extracted
        assert len(llm.calls) == 0  # no point calling the LLM with no text
    finally:
        db.close()


def test_ledger_summary_counts_by_source_type_and_status():
    db = SessionLocal()
    try:
        record_item(
            db,
            source_type="cv",
            source_ref="a.pdf",
            extracted={"purpose": "x", "stack": ["Python"], "role": "y", "highlights": []},
            source_urls=[],
        )
        record_item(
            db, source_type="cv", source_ref="b.pdf", extracted={"parse_error": "bad"}, source_urls=[]
        )
        summary = ledger_summary(db)
        assert summary["cv"]["analyzed"] >= 1
        assert summary["cv"]["needs_review"] >= 1
    finally:
        db.close()


def test_reingesting_a_source_updates_its_record_instead_of_duplicating():
    from app.models import ProfileItem

    db = SessionLocal()
    try:
        ref = "owner/upsert-repo"
        good = {"purpose": "v1", "stack": ["Python"], "role": "author", "highlights": []}
        first = record_item(db, source_type="github_repo", source_ref=ref, extracted=good, source_urls=[])
        # A copy left over from before re-ingestion updated in place.
        db.add(ProfileItem(source_type="github_repo", source_ref=ref, extracted=good, status=ProfileItemStatus.ANALYZED))
        db.commit()

        newer = {**good, "purpose": "v2"}
        second = record_item(db, source_type="github_repo", source_ref=ref, extracted=newer, source_urls=[])
        rows = db.query(ProfileItem).filter(ProfileItem.source_ref == ref).all()
        assert len(rows) == 1 and rows[0].id == second.id
        assert rows[0].extracted["purpose"] == "v2"

        # A failed re-run never replaces the good record.
        failed = {"purpose": "unknown", "stack": [], "role": "unknown", "highlights": [], "parse_error": "boom"}
        kept = record_item(db, source_type="github_repo", source_ref=ref, extracted=failed, source_urls=[])
        assert kept.status == ProfileItemStatus.ANALYZED and kept.extracted["purpose"] == "v2"
        assert db.query(ProfileItem).filter(ProfileItem.source_ref == ref).count() == 1
        assert first.id  # first record existed before the fold
    finally:
        db.query(ProfileItem).filter(ProfileItem.source_ref == "owner/upsert-repo").delete()
        db.commit()
        db.close()
